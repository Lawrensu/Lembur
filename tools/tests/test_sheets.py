import datetime as dt
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import schema  # noqa: E402
import sheets  # noqa: E402

try:
    import openpyxl
except ImportError:
    openpyxl = None


def blank(sheet):
    return {name: "" for name in schema.LEDGER_SHEETS[sheet]}


def sample_book():
    sale = blank("transactions")
    sale.update(
        txn_id="T0001", date="2026-10-05", type="income", category="sales",
        description="Kuih tray", product_id="P01", quantity="3",
        amount_rm="45.5", channel="whatsapp", payment_method="cash",
    )
    expense = blank("transactions")
    expense.update(
        txn_id="T0002", date="2026-10-06", type="expense", category="materials",
        description="Flour", amount_rm="12", payment_method="bank_transfer",
        notes="receipt, blurry",
    )
    bank = blank("bank_lines")
    bank.update(
        date="2026-10-06", description="TRANSFER OUT", type="expense",
        amount_rm="12.00", bank_reference="B001",
    )
    hours = blank("hours")
    hours.update(
        entry_id="H0001", date="2026-10-05", person="person_1",
        activity="production", product_id="P01", quantity="3", minutes="90",
        source="derived", assumption_ref="products.P01.minutes_per_unit",
    )
    return {"transactions": [sale, expense], "bank_lines": [bank], "hours": [hours]}


class FormatTests(unittest.TestCase):
    def test_csv_and_auto(self):
        self.assertEqual(sheets.resolve_format("csv"), "csv")
        expected = "xlsx" if openpyxl else "csv"
        self.assertEqual(sheets.resolve_format("auto"), expected)

    def test_unknown_format(self):
        with self.assertRaises(ValueError):
            sheets.resolve_format("ods")

    @unittest.skipIf(openpyxl, "openpyxl is installed")
    def test_xlsx_without_openpyxl(self):
        with self.assertRaises(RuntimeError):
            sheets.resolve_format("xlsx")


class CsvTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = os.path.join(self.tmp.name, "workspace", "ledger.xlsx")

    def test_round_trip(self):
        book = sample_book()
        sheets.save_book(self.path, schema.LEDGER_SHEETS, book, "csv")
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "csv")
        self.assertEqual(list(loaded), list(schema.LEDGER_SHEETS))
        for sheet, columns in schema.LEDGER_SHEETS.items():
            for row in loaded[sheet]:
                self.assertEqual(list(row), columns)
        self.assertEqual(loaded["transactions"][0]["product_id"], "P01")
        self.assertEqual(loaded["transactions"][0]["notes"], "")
        self.assertEqual(loaded["transactions"][1]["notes"], "receipt, blurry")
        self.assertEqual(loaded["hours"], book["hours"])

    def test_file_name_and_money_two_decimals(self):
        sheets.save_book(self.path, schema.LEDGER_SHEETS, sample_book(), "csv")
        csv_file = os.path.join(self.tmp.name, "workspace", "ledger_transactions.csv")
        self.assertTrue(os.path.isfile(csv_file))
        with open(csv_file, encoding="utf-8", newline="") as handle:
            lines = handle.read().splitlines()
        self.assertEqual(lines[0].split(",")[:2], ["txn_id", "date"])
        self.assertIn(",45.50,", lines[1])
        self.assertIn(",12.00,", lines[2])
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "csv")
        self.assertEqual(loaded["transactions"][0]["amount_rm"], "45.50")

    def test_missing_file(self):
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "csv")
        self.assertEqual(loaded, {sheet: [] for sheet in schema.LEDGER_SHEETS})

    def test_empty_rows_skipped(self):
        book = sample_book()
        book["transactions"].insert(1, blank("transactions"))
        sheets.save_book(self.path, schema.LEDGER_SHEETS, book, "csv")
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "csv")
        self.assertEqual(len(loaded["transactions"]), 2)

    def test_missing_sheet_in_book_gets_header_only(self):
        sheets.save_book(self.path, schema.LEDGER_SHEETS, {}, "csv")
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "csv")
        self.assertEqual(loaded, {sheet: [] for sheet in schema.LEDGER_SHEETS})
        csv_file = os.path.join(self.tmp.name, "workspace", "ledger_hours.csv")
        with open(csv_file, encoding="utf-8", newline="") as handle:
            self.assertEqual(handle.read().strip(), ",".join(schema.LEDGER_SHEETS["hours"]))

    def test_extra_and_missing_columns(self):
        csv_file = os.path.join(self.tmp.name, "ledger_bank_lines.csv")
        with open(csv_file, "w", encoding="utf-8-sig", newline="") as handle:
            handle.write("bank_reference,junk,amount_rm,date\r\nB9, x ,5,2026-10-01\r\n")
        path = os.path.join(self.tmp.name, "ledger.xlsx")
        loaded = sheets.load_book(path, {"bank_lines": schema.LEDGER_SHEETS["bank_lines"]}, "csv")
        self.assertEqual(
            loaded["bank_lines"],
            [{"date": "2026-10-01", "description": "", "type": "",
              "amount_rm": "5", "bank_reference": "B9"}],
        )

    def test_empty_money_stays_empty(self):
        book = sample_book()
        book["transactions"][0]["amount_rm"] = ""
        sheets.save_book(self.path, schema.LEDGER_SHEETS, book, "csv")
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "csv")
        self.assertEqual(loaded["transactions"][0]["amount_rm"], "")


@unittest.skipUnless(openpyxl, "openpyxl is not installed")
class XlsxTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = os.path.join(self.tmp.name, "workspace", "ledger.xlsx")

    def test_round_trip(self):
        book = sample_book()
        sheets.save_book(self.path, schema.LEDGER_SHEETS, book, "xlsx")
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "xlsx")
        self.assertEqual(list(loaded), list(schema.LEDGER_SHEETS))
        self.assertEqual(loaded["transactions"][0]["date"], "2026-10-05")
        self.assertEqual(loaded["transactions"][0]["amount_rm"], "45.5")
        self.assertEqual(loaded["transactions"][1]["amount_rm"], "12")
        self.assertEqual(loaded["transactions"][0]["quantity"], "3")
        self.assertEqual(loaded["transactions"][0]["notes"], "")
        self.assertEqual(loaded["hours"], book["hours"])
        for sheet, columns in schema.LEDGER_SHEETS.items():
            for row in loaded[sheet]:
                self.assertEqual(list(row), columns)

    def test_cell_types_and_formats(self):
        sheets.save_book(self.path, schema.LEDGER_SHEETS, sample_book(), "xlsx")
        ws = openpyxl.load_workbook(self.path)["transactions"]
        columns = schema.LEDGER_SHEETS["transactions"]
        amount = ws.cell(row=2, column=columns.index("amount_rm") + 1)
        self.assertEqual(amount.value, 45.5)
        self.assertIsInstance(amount.value, float)
        self.assertEqual(amount.number_format, "0.00")
        quantity = ws.cell(row=2, column=columns.index("quantity") + 1)
        self.assertEqual(quantity.value, 3)
        self.assertIsInstance(quantity.value, int)
        self.assertEqual(ws.cell(row=2, column=columns.index("date") + 1).value, "2026-10-05")
        self.assertIsNone(ws.cell(row=2, column=columns.index("notes") + 1).value)
        self.assertIsNone(ws.cell(row=3, column=columns.index("product_id") + 1).value)

    def test_missing_file_and_sheet(self):
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "xlsx")
        self.assertEqual(loaded, {sheet: [] for sheet in schema.LEDGER_SHEETS})
        os.makedirs(os.path.dirname(self.path))
        workbook = openpyxl.Workbook()
        workbook.active.title = "bank_lines"
        workbook.save(self.path)
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "xlsx")
        self.assertEqual(loaded, {sheet: [] for sheet in schema.LEDGER_SHEETS})

    def test_empty_rows_skipped(self):
        book = sample_book()
        book["transactions"].insert(1, blank("transactions"))
        sheets.save_book(self.path, schema.LEDGER_SHEETS, book, "xlsx")
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "xlsx")
        self.assertEqual(len(loaded["transactions"]), 2)

    def test_hand_made_cells(self):
        os.makedirs(os.path.dirname(self.path))
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = "bank_lines"
        ws.append(["bank_reference", "extra", "amount_rm", "date", "type"])
        ws.append(["B1", "ignored", 15.0, dt.datetime(2026, 10, 7, 0, 0), "income"])
        ws.append(["B2", None, 23.5, dt.date(2026, 10, 8), "expense"])
        ws.append([None, None, None, None, None])
        ws.append(["  B3  ", None, 7, "2026-10-09", None])
        workbook.save(self.path)
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "xlsx")["bank_lines"]
        self.assertEqual([row["amount_rm"] for row in loaded], ["15", "23.5", "7"])
        self.assertEqual([row["date"] for row in loaded], ["2026-10-07", "2026-10-08", "2026-10-09"])
        self.assertEqual([row["bank_reference"] for row in loaded], ["B1", "B2", "B3"])
        self.assertEqual(loaded[0]["description"], "")
        self.assertEqual(list(loaded[0]), schema.LEDGER_SHEETS["bank_lines"])
        self.assertNotIn("extra", loaded[0])

    def test_existing_extra_sheet_survives(self):
        os.makedirs(os.path.dirname(self.path))
        workbook = openpyxl.Workbook()
        notes = workbook.active
        notes.title = "student_notes"
        notes["A1"] = "keep me"
        notes.column_dimensions["A"].width = 33
        old = workbook.create_sheet("hours")
        old.append(["entry_id", "stale"])
        old.append(["H9", "old row"])
        workbook.save(self.path)

        sheets.save_book(self.path, schema.LEDGER_SHEETS, sample_book(), "xlsx")

        reopened = openpyxl.load_workbook(self.path)
        self.assertEqual(reopened["student_notes"]["A1"].value, "keep me")
        self.assertEqual(reopened["student_notes"].column_dimensions["A"].width, 33)
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "xlsx")
        self.assertEqual([row["entry_id"] for row in loaded["hours"]], ["H0001"])
        self.assertEqual(len(loaded["transactions"]), 2)

    def test_second_save_replaces_rows(self):
        sheets.save_book(self.path, schema.LEDGER_SHEETS, sample_book(), "xlsx")
        smaller = sample_book()
        smaller["transactions"] = smaller["transactions"][:1]
        sheets.save_book(self.path, schema.LEDGER_SHEETS, smaller, "xlsx")
        loaded = sheets.load_book(self.path, schema.LEDGER_SHEETS, "xlsx")
        self.assertEqual(len(loaded["transactions"]), 1)

    def test_default_sheet_removed(self):
        sheets.save_book(self.path, schema.LEDGER_SHEETS, sample_book(), "xlsx")
        names = openpyxl.load_workbook(self.path).sheetnames
        self.assertEqual(names, list(schema.LEDGER_SHEETS))


if __name__ == "__main__":
    unittest.main()

"""Read and write the workbook files as plain strings.

A book is a dict of sheet name to a list of rows; a row is a dict of column
name to string, with "" for an empty cell. Uses .xlsx through openpyxl when it
is installed, otherwise one CSV file per sheet (docs/SYSTEM_DESIGN.md 3.1).
"""

from __future__ import annotations

import csv
import datetime as dt
import os
from decimal import Decimal, InvalidOperation

from conventions import is_empty, rm, text, to_decimal, to_sen
from schema import MONEY_COLUMNS, NUMBER_COLUMNS


def _import_openpyxl():
    try:
        import openpyxl
    except ImportError:
        return None
    return openpyxl


def resolve_format(fmt="auto") -> str:
    """Return "xlsx" or "csv". "auto" picks xlsx when openpyxl is installed."""
    if fmt == "auto":
        return "xlsx" if _import_openpyxl() else "csv"
    if fmt == "xlsx":
        if _import_openpyxl() is None:
            raise RuntimeError(
                "Format xlsx needs the openpyxl package (pip install openpyxl); "
                "use fmt='csv' or fmt='auto' to fall back to CSV files."
            )
        return "xlsx"
    if fmt == "csv":
        return "csv"
    raise ValueError("Unknown format %r; use 'auto', 'xlsx' or 'csv'." % (fmt,))


def _csv_path(path, sheet):
    return "%s_%s.csv" % (os.path.splitext(str(path))[0], sheet)


def _cell_text(value) -> str:
    """One spreadsheet cell as the string the other tools see."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return format(Decimal(str(value)), "f")
    return text(value)


def _shape(raw_rows, columns):
    """Keep the wanted columns in order and drop rows that are entirely empty."""
    rows = []
    for raw in raw_rows:
        row = {name: _cell_text(raw.get(name)) for name in columns}
        if any(row.values()):
            rows.append(row)
    return rows


def _read_xlsx_sheet(workbook, sheet, columns):
    if sheet not in workbook.sheetnames:
        return []
    lines = workbook[sheet].iter_rows(values_only=True)
    header = next(lines, None)
    if header is None:
        return []
    names = [text(cell) for cell in header]
    raw_rows = []
    for line in lines:
        raw = {}
        for name, cell in zip(names, line):
            if name and name not in raw:
                raw[name] = cell
        raw_rows.append(raw)
    return _shape(raw_rows, columns)


def _read_csv_sheet(path, columns):
    if not os.path.isfile(path):
        return []
    with open(path, encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        raw_rows = [
            {text(key): value for key, value in record.items() if key is not None}
            for record in reader
        ]
    return _shape(raw_rows, columns)


def load_book(path, sheet_columns, fmt="auto") -> dict:
    """Read every sheet in sheet_columns into a book of string rows.

    Missing files and missing sheets give an empty list. Extra columns are
    ignored and missing ones come back as "".
    """
    if resolve_format(fmt) == "csv":
        return {
            sheet: _read_csv_sheet(_csv_path(path, sheet), columns)
            for sheet, columns in sheet_columns.items()
        }
    if not os.path.isfile(path):
        return {sheet: [] for sheet in sheet_columns}
    openpyxl = _import_openpyxl()
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        return {
            sheet: _read_xlsx_sheet(workbook, sheet, columns)
            for sheet, columns in sheet_columns.items()
        }
    finally:
        workbook.close()


def _xlsx_value(column, value):
    """The Python value to put in a cell; None leaves the cell empty."""
    if is_empty(value):
        return None
    try:
        if column in MONEY_COLUMNS:
            return to_sen(value) / 100
        if column in NUMBER_COLUMNS:
            number = to_decimal(value)
            if number == number.to_integral_value():
                return int(number)
            return float(number)
    except InvalidOperation:
        pass  # not a number (for example a hand typed note): keep it as text
    return text(value)


def _csv_value(column, value):
    if column in MONEY_COLUMNS and not is_empty(value):
        try:
            return rm(to_sen(value))
        except InvalidOperation:
            pass
    return text(value)


def _write_xlsx(path, sheet_columns, book):
    openpyxl = _import_openpyxl()
    if os.path.isfile(path):
        workbook = openpyxl.load_workbook(path)
    else:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
    for sheet, columns in sheet_columns.items():
        if sheet in workbook.sheetnames:
            ws = workbook[sheet]
        else:
            ws = workbook.create_sheet(sheet)
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for index, name in enumerate(columns, start=1):
            ws.cell(row=1, column=index, value=name)
        for line, row in enumerate(book.get(sheet, []), start=2):
            for index, name in enumerate(columns, start=1):
                cell = ws.cell(
                    row=line, column=index, value=_xlsx_value(name, row.get(name))
                )
                if name in MONEY_COLUMNS and cell.value is not None:
                    cell.number_format = "0.00"
    workbook.save(path)


def _write_csv(path, sheet_columns, book):
    for sheet, columns in sheet_columns.items():
        with open(_csv_path(path, sheet), "w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(columns)
            for row in book.get(sheet, []):
                writer.writerow([_csv_value(name, row.get(name)) for name in columns])


def save_book(path, sheet_columns, book, fmt="auto") -> None:
    """Write every sheet in sheet_columns from the book.

    In xlsx mode an existing file keeps its other sheets and column widths;
    the named sheets are rewritten below the header row.
    """
    mode = resolve_format(fmt)
    folder = os.path.dirname(os.path.abspath(path))
    os.makedirs(folder, exist_ok=True)
    if mode == "csv":
        _write_csv(path, sheet_columns, book)
    else:
        _write_xlsx(path, sheet_columns, book)

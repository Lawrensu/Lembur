"""Checks whether this host can run deterministic Python scripts, and which Python and libraries it has.
Run from the repo root with: python tools/check_runtime.py"""

import csv
import datetime
import os
import platform
import sys
import tempfile
from decimal import Decimal, ROUND_HALF_UP

OPENPYXL = None


def report(name, ok, detail):
    print("CHECK %s: %s %s" % (name, "OK" if ok else "FAIL", detail))
    return ok


def run_check(name, func):
    try:
        ok, detail = func()
    except Exception as exc:
        ok, detail = False, "%s: %s" % (type(exc).__name__, exc)
    return report(name, ok, detail)


def check_python():
    return True, "version %s, executable %s" % (
        sys.version.replace("\n", " "),
        sys.executable,
    )


def check_platform():
    return True, "%s, cwd %s" % (platform.platform(), os.getcwd())


def check_decimal():
    raw = Decimal(6800) * 60 / Decimal(1340)
    rate_sen = raw.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    half = Decimal("304.5").quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    ok = rate_sen == Decimal(304) and half == Decimal(305)
    ringgit = (rate_sen / Decimal(100)).quantize(Decimal("0.01"))
    detail = (
        "raw %s sen per hour, rounded %s sen (RM%s); Decimal('304.5') rounded %s"
        % (raw, rate_sen, ringgit, half)
    )
    return ok, detail


def check_iso_week():
    year, week, weekday = datetime.date(2026, 10, 10).isocalendar()
    monday = datetime.date.fromisocalendar(2026, 41, 1)
    ok = (year, week) == (2026, 41) and monday == datetime.date(2026, 10, 5)
    detail = "2026-10-10 is ISO %d-W%02d-%d; fromisocalendar(2026, 41, 1) is %s" % (
        year,
        week,
        weekday,
        monday.isoformat(),
    )
    return ok, detail


CSV_ROWS = [
    ["date", "description", "amount"],
    ["2026-10-01", "Nasi lemak", "23.50"],
    ["2026-10-02", "", "12.00"],
]


def check_csv():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "sample.csv")
        with open(path, "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows(CSV_ROWS)
        with open(path, "r", newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
    ok = rows == CSV_ROWS
    if ok:
        return ok, "header and 2 rows read back identical (amount 23.50, empty cell kept)"
    return ok, "read back %r, expected %r" % (rows, CSV_ROWS)


def check_openpyxl():
    global OPENPYXL
    try:
        import openpyxl
    except ImportError:
        return False, "openpyxl is not installed"
    OPENPYXL = openpyxl
    return True, "openpyxl %s" % getattr(openpyxl, "__version__", "version unknown")


XLSX_HEADER = ["description", "amount", "note"]
XLSX_ROW = ["Nasi lemak", 23.5, None]


def check_xlsx():
    if OPENPYXL is None:
        return False, "skipped, openpyxl not installed"
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "sample.xlsx")
        wb = OPENPYXL.Workbook()
        ws = wb.active
        ws.title = "transactions"
        ws.append(XLSX_HEADER)
        ws.append(XLSX_ROW)
        wb.save(path)
        wb2 = OPENPYXL.load_workbook(path)
        ws2 = wb2["transactions"]
        rows = [list(r) for r in ws2.iter_rows(values_only=True)]
    ok = rows == [XLSX_HEADER, XLSX_ROW]
    if ok:
        return ok, "sheet transactions saved and reopened; text, 23.5 and empty cell match"
    return ok, "read back %r, expected %r" % (rows, [XLSX_HEADER, XLSX_ROW])


def check_write_access():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "write_test.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write("ok")
        created = os.path.isfile(path)
        os.remove(path)
        removed = not os.path.exists(path)
    ok = created and removed
    return ok, "created and deleted a file in temp directory %s" % tmp


def main():
    try:
        if hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(errors="backslashreplace")
    except Exception:
        pass

    results = {}
    results["python"] = run_check("python", check_python)
    results["platform"] = run_check("platform", check_platform)
    results["decimal_rounding"] = run_check("decimal_rounding", check_decimal)
    results["iso_week"] = run_check("iso_week", check_iso_week)
    results["csv_round_trip"] = run_check("csv_round_trip", check_csv)
    results["openpyxl_import"] = run_check("openpyxl_import", check_openpyxl)
    results["xlsx_round_trip"] = run_check("xlsx_round_trip", check_xlsx)
    results["write_access"] = run_check("write_access", check_write_access)

    total = len(results)
    ok_count = sum(1 for passed in results.values() if passed)
    print("SUMMARY: %d of %d checks OK" % (ok_count, total))

    if results["openpyxl_import"] and results["xlsx_round_trip"]:
        mode = "xlsx"
    elif results["csv_round_trip"]:
        mode = "csv"
    else:
        mode = "none"
    print("SCRIPT MODE: %s" % mode)


if __name__ == "__main__":
    main()

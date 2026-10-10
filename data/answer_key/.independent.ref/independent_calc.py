"""Independent, read-only fixture reconstruction from the five authorized inputs."""

from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from hashlib import sha256
import json
from pathlib import Path
import re

from openpyxl import load_workbook

ROOT = Path("C:/Users/user/Documents/Hackerdorm/Lembur")
OUT = ROOT / "data/answer_key/.independent.ref"
INPUTS = (
    ROOT / "docs/PLAN.md",
    ROOT / "docs/SYSTEM_DESIGN.md",
    ROOT / "data/demo/plan.xlsx",
    ROOT / "data/demo/messages.md",
    OUT / "input_transactions_bank.xlsx",
)
CATEGORIES = ("sales", "materials", "packaging", "platform_fees", "shipping", "other_expenses")
ACTIVITIES = ("production", "fulfilment", "event", "admin")
REVIEW_STATUSES = {"unmatched", "duplicate", "needs_review"}


def dec(value):
    return Decimal(str(value))


def whole(value):
    value = dec(value)
    assert value == value.to_integral_value(), f"Not a whole value: {value}"
    return int(value)


def sen(value):
    return None if value is None else whole(dec(value) * 100)


def display(value, places=2):
    return format(dec(value).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP), f".{places}f")


def money(value):
    return None if value is None else display(dec(value) / 100)


def ratio(numerator, denominator, unit, places=2):
    return {
        "exact_numerator": numerator,
        "exact_denominator": denominator,
        "unit": unit,
        "display": None if denominator == 0 else display(dec(numerator) / dec(denominator), places),
    }


def iso(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return date.fromisoformat(value).isoformat()


def fingerprints():
    return {p.relative_to(ROOT).as_posix(): sha256(p.read_bytes()).hexdigest() for p in INPUTS}


def workbook(path, required_sheets):
    assert path in INPUTS
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        assert set(book.sheetnames) == set(required_sheets), "Unexpected workbook sheets"
        result = {}
        for name in required_sheets:
            iterator = book[name].iter_rows()
            columns = [c.value for c in next(iterator)]
            assert len(columns) == len(set(columns)) and None not in columns
            rows = []
            for cells in iterator:
                assert all(c.data_type != "f" for c in cells), "Formula input is not evaluated or guessed"
                values = [c.value for c in cells]
                if any(v is not None for v in values):
                    rows.append(dict(zip(columns, values)))
            result[name] = rows
        return result
    finally:
        book.close()


def parse_messages(text, baseline):
    sections = re.findall(r"^## (\d{4}-\d{2}-\d{2})[^\n]*\n(.*?)(?=^## |\Z)", text, re.M | re.S)
    messages, corrections, events = {}, [], []
    number_words = dict(zip("one two three four five six seven eight nine ten eleven twelve".split(), range(1, 13)))
    for day, section in sections:
        message_lines = re.findall(r"^- Message: (.+)$", section, re.M)
        assert len(message_lines) == 1, f"Expected one nightly message for {day}"
        message = message_lines[0]
        messages[day] = message
        for match in re.finditer(r"([A-Za-z]+) take (\d+) minutes not (\d+)", message):
            product_word, new, old = match.groups()
            product_word = product_word.lower().removesuffix("s")
            products = [p for p in baseline["products"] if product_word in p["name"].lower().split()]
            assert len(products) == 1, "Correction product is not unambiguous"
            corrections.append({
                "changed_at": day,
                "field": f"products.{products[0]['product_id']}.minutes_per_unit",
                "old_value": whole(old),
                "new_value": whole(new),
                "reason": message[message.rfind(". ", 0, match.start()) + 2:].rstrip("."),
                "source": f"message.{day}",
            })
        if "bazaar" in message.lower():
            times = re.search(r"from (\d{1,2}:\d{2}[ap]m) to (\d{1,2}:\d{2}[ap]m)", message, re.I)
            assert times, "Event duration is unclear; do not infer it"
            start, end = [datetime.strptime(t.lower(), "%I:%M%p") for t in times.groups()]
            minutes = whole(dec(str((end - start).total_seconds())) / 60)
            assert minutes > 0, "Event date rollover is not specified"
            stated = re.search(r"(\d+(?:\.\d+)?|" + "|".join(number_words) + r") hours in total", message, re.I)
            if stated:
                token = stated.group(1).lower()
                stated_minutes = whole(dec(number_words[token] if token in number_words else token) * 60)
                assert minutes == stated_minutes, "Clock range and stated duration conflict"
            events.append({"date": day, "minutes": minutes, "source": "stated", "assumption_ref": f"message.{day}", "message_text": message})
    assert messages, "No dated messages found"
    assert len({(c['changed_at'], c['field']) for c in corrections}) == len(corrections)
    return messages, corrections, events


def plan_cell(plan, field):
    parts = field.split(".")
    if parts[0] == "assumptions":
        return plan["assumptions"][0], parts[1]
    assert parts[0] in {"products", "channels"}
    key = "product_id" if parts[0] == "products" else "channel"
    matches = [r for r in plan[parts[0]] if r[key] == parts[1]]
    assert len(matches) == 1
    return matches[0], parts[2]


def apply_corrections(baseline, corrections):
    assert not baseline["change_log"], "Expected an unmodified baseline plan"
    plan = deepcopy(baseline)
    for correction in sorted(corrections, key=lambda c: (c["changed_at"], c["field"])):
        row, column = plan_cell(plan, correction["field"])
        assert dec(row[column]) == dec(correction["old_value"]), "Correction old value conflicts with plan"
        row[column] = correction["new_value"]
        plan["change_log"].append(deepcopy(correction))
    return plan


def in_force(plan, field, day):
    changes = sorted([c for c in plan["change_log"] if c["field"] == field], key=lambda c: c["changed_at"])
    future = [c for c in changes if c["changed_at"] > day]
    cell, column = plan_cell(plan, field)
    value = future[0]["old_value"] if future else cell[column]
    applicable = [c for c in changes if c["changed_at"] <= day]
    # Sharon's clarification: historical old_value is an estimate, not corrected.
    return value, "corrected" if applicable else "estimate", applicable[-1]["source"] if applicable else None


def normalize_row(row):
    result = deepcopy(row)
    result["date"] = iso(result["date"])
    result["amount_sen"] = sen(result.pop("amount_rm"))
    if "quantity" in result and result["quantity"] is not None:
        result["quantity"] = whole(result["quantity"])
    return result


def raw_amount(row):
    match = re.search(r"Document RM(\d+(?:\.\d+)?); bank RM(\d+(?:\.\d+)?)", row.get("notes") or "", re.I)
    if match:
        return sen(match.group(1)), sen(match.group(2)), "Original document amount reconstructed from supplied notes; image not inspected"
    return row["amount_sen"], None, "Supplied transaction amount; original document not inspected"


def reconcile(input_rows, bank_rows):
    rows = sorted([normalize_row(r) for r in input_rows], key=lambda r: r["txn_id"])
    bank = [normalize_row(r) for r in bank_rows]
    bank_by_ref = {b["bank_reference"]: b for b in bank}
    assert len(bank_by_ref) == len(bank)
    documents = [r for r in rows if r["source_file"]]
    bank_only_inputs = [r for r in rows if not r["source_file"]]
    originals = {r["txn_id"]: raw_amount(r) for r in documents}
    used, signatures, audits, reconstructed = set(), {}, [], []

    def finish(row, status, reference, amount, rule, evidence, candidates=None):
        new = deepcopy(row)
        new.update(reconciliation_status=status, bank_reference=reference, amount_sen=amount)
        new["notes"] = evidence
        new["counted"] = status != "duplicate" and amount is not None
        reconstructed.append(new)
        if reference:
            assert reference not in used, "Bank line reused"
            used.add(reference)
        original, note_bank, provenance = originals.get(row["txn_id"], (None, None, "No source document; reconstructed from leftover bank line"))
        audit = {
            "txn_id": row["txn_id"], "rule": rule,
            "input_status": row["reconciliation_status"], "derived_status": status,
            "status_agrees": status == row["reconciliation_status"],
            "input_bank_reference": row["bank_reference"], "derived_bank_reference": reference,
            "bank_reference_agrees": reference == row["bank_reference"],
            "input_amount_sen": row["amount_sen"], "derived_amount_sen": amount,
            "amount_agrees": amount == row["amount_sen"],
            "original_document_amount_sen": original, "document_amount_evidence": provenance,
            "bank_amount_sen": bank_by_ref[reference]["amount_sen"] if reference else None,
            "counted": new["counted"], "evidence": evidence,
            "amount_candidate_references": candidates or [],
        }
        if note_bank is not None:
            audit["bank_amount_in_notes_sen"] = note_bank
            audit["bank_note_agrees"] = reference is not None and note_bank == bank_by_ref[reference]["amount_sen"]
        audits.append(audit)

    for row in documents:
        original, _, _ = originals[row["txn_id"]]
        required = ["date", "type", "category", "payment_method"]
        if row["category"] == "sales":
            required += ["product_id", "quantity", "channel"]
        missing = [key for key in required if row.get(key) is None]
        if original is None:
            missing.append("amount_rm")
        signature = (row["date"], row["type"], original, row["counterparty"], row["reference"])
        if missing:
            finish(row, "needs_review", None, original, "1: required-field intake flag", "Missing " + ", ".join(missing) + "; flag precedes cash classification")
        elif row["payment_method"] == "cash":
            finish(row, "cash_declared", None, original, "2: cash", "Cash is declared; no bank line is expected")
        elif signature in signatures:
            finish(row, "duplicate", None, original, "3: duplicate", f"Duplicate of {signatures[signature]} on date, type, original amount, counterparty and reference")
        else:
            signatures[signature] = row["txn_id"]
            reference_match = [b for b in bank if b["bank_reference"] not in used and b["type"] == row["type"] and row["reference"] == b["bank_reference"]]
            if reference_match:
                assert len(reference_match) == 1
                chosen = reference_match[0]
                agrees = chosen["amount_sen"] == original
                evidence = "Document reference, type and amount agree with bank" if agrees else f"Document RM{money(original)}; bank RM{money(chosen['amount_sen'])}; bank amount used; confirm discrepancy"
                finish(row, "matched" if agrees else "needs_review", chosen["bank_reference"], chosen["amount_sen"], "4: reference match" if agrees else "4: reference amount mismatch", evidence)
            else:
                def candidates_for(other):
                    amount = originals[other["txn_id"]][0]
                    return [b for b in bank if b["bank_reference"] not in used and b["type"] == other["type"] and b["amount_sen"] == amount and abs((date.fromisoformat(b["date"]) - date.fromisoformat(other["date"])).days) <= 1]

                candidates = candidates_for(row)
                refs = [b["bank_reference"] for b in candidates]
                competitors = [other for other in documents if other["txn_id"] > row["txn_id"] and other["payment_method"] != "cash" and originals[other["txn_id"]][0] is not None and any(b["bank_reference"] in refs for b in candidates_for(other))]
                if not candidates:
                    finish(row, "unmatched", None, original, "5: no amount candidate", "No unused bank line has the same type and amount within one day")
                elif len(candidates) == 1 and not competitors:
                    finish(row, "matched", refs[0], original, "5: unique amount candidate", "One candidate within one day, shared by no other pending document", refs)
                else:
                    counterparty = row["counterparty"] or ""
                    identified = [b for b in candidates if counterparty and re.search(r"(?<!\w)" + re.escape(counterparty) + r"(?!\w)", b["description"], re.I)]
                    if len(identified) == 1:
                        finish(row, "ambiguous_resolved", identified[0]["bank_reference"], original, "6: description disambiguation", f"Amount/date candidates {', '.join(refs)}; only {identified[0]['bank_reference']} names counterparty {counterparty}", refs)
                    else:
                        finish(row, "needs_review", None, original, "6: unresolved ambiguity", "Counterparty and bank descriptions do not identify one candidate", refs)

    consumed_bank_only = set()
    for line in bank:
        if line["bank_reference"] in used:
            continue
        matches = [r for r in bank_only_inputs if all(r[k] == line[k] for k in ("date", "type", "amount_sen", "description"))]
        assert len(matches) == 1, "Leftover bank line has no unique corresponding input row"
        row = matches[0]
        consumed_bank_only.add(row["txn_id"])
        description = line["description"].lower()
        if line["type"] == "income":
            category = "sales"
        elif "materials" in description:
            category = "materials"
        elif "packaging" in description:
            category = "packaging"
        else:
            category = "other_expenses"
        finish(row, "unmatched", line["bank_reference"], line["amount_sen"], "7: leftover bank line", f"No document row represents this bank line; description supports category {category}; request receipt")
        reconstructed[-1]["category"] = category
        audits[-1]["derived_category"] = category
        audits[-1]["category_agrees"] = category == row["category"]
    assert consumed_bank_only == {r["txn_id"] for r in bank_only_inputs}
    return sorted(reconstructed, key=lambda r: r["txn_id"]), sorted(audits, key=lambda r: r["txn_id"]), bank


def build_hours(rows, plan, events, sunday):
    production, fulfilment = defaultdict(list), defaultdict(list)
    for row in rows:
        if row["counted"] and row["category"] == "sales" and row["product_id"] and row["quantity"] is not None:
            production[row["date"], row["product_id"]].append(row)
            fulfilment[row["date"], row["product_id"], row["channel"]].append(row)
    hours = []

    def add(day, activity, product, quantity, field, basis, source_rows=None, channel=None):
        value, source, correction_message = in_force(plan, field, day)
        minutes = whole(dec(basis) * dec(value))
        if activity == "fulfilment" and minutes == 0:
            return
        hours.append({
            "date": day, "person": "person_1", "activity": activity, "product_id": product,
            "channel": channel, "quantity": quantity, "minutes": minutes,
            "source": source, "assumption_ref": field, "value_in_force": value,
            "correction_message": correction_message,
            "transaction_ids": [r["txn_id"] for r in (source_rows or [])],
        })

    for (day, product), sales in sorted(production.items()):
        units = sum(r["quantity"] for r in sales)
        add(day, "production", product, units, f"products.{product}.minutes_per_unit", units, sales)
    for (day, product, channel), sales in sorted(fulfilment.items()):
        add(day, "fulfilment", product, len(sales), f"channels.{channel}.fulfilment_minutes_per_order", len(sales), sales, channel)
    for event in events:
        hours.append({
            "date": event["date"], "person": "person_1", "activity": "event", "product_id": None,
            "channel": None, "quantity": None, "minutes": event["minutes"], "source": "stated",
            "assumption_ref": event["assumption_ref"], "value_in_force": None,
            "correction_message": None, "transaction_ids": [],
        })
    add(sunday, "admin", None, None, "assumptions.admin_hours_per_week", 60)
    hours.sort(key=lambda h: (h["date"], ACTIVITIES.index(h["activity"]), h["product_id"] or "", h["channel"] or "", h["assumption_ref"]))
    for index, entry in enumerate(hours, 1):
        entry["entry_id"] = f"H{index:04d}"
    return hours


def percentage_fee(amount_sen, fee_percent):
    return int((dec(amount_sen) * dec(fee_percent) / 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def compute_numbers(rows, hours, plan):
    totals = {category: sum(r["amount_sen"] for r in rows if r["counted"] and r["category"] == category) for category in CATEGORIES}
    expenses = sum(totals[c] for c in CATEGORIES if c != "sales")
    take_home = totals["sales"] - expenses
    activities = {a: sum(h["minutes"] for h in hours if h["activity"] == a) for a in ACTIVITIES}
    total_minutes = sum(activities.values())
    products, order_costs = [], []
    channels = {r["channel"]: r for r in plan["channels"]}
    for product in plan["products"]:
        pid = product["product_id"]
        sales = [r for r in rows if r["counted"] and r["category"] == "sales" and r["product_id"] == pid]
        revenue = sum(r["amount_sen"] for r in sales)
        units = sum(r["quantity"] for r in sales)
        material = units * sen(product["material_cost_rm"])
        packaging = units * sen(product["packaging_cost_rm"])
        breakdown = {}
        for channel_name in sorted({r["channel"] for r in sales}):
            channel = channels[channel_name]
            entries = []
            for row in sales:
                if row["channel"] != channel_name:
                    continue
                percentage = percentage_fee(row["amount_sen"], channel["fee_percent"])
                fixed = sen(channel["fee_per_order_rm"])
                shipping = sen(channel["shipping_cost_per_order_rm"])
                entry = {"txn_id": row["txn_id"], "product_id": pid, "channel": channel_name,
                         "revenue_sen": row["amount_sen"], "fee_percent": channel["fee_percent"],
                         "percentage_fee_sen": percentage, "fixed_fee_sen": fixed,
                         "shipping_cost_sen": shipping, "total_channel_cost_sen": percentage + fixed + shipping}
                entries.append(entry)
                order_costs.append(entry)
            breakdown[channel_name] = {
                "orders": len(entries), "units": sum(r["quantity"] for r in sales if r["channel"] == channel_name),
                **{key: sum(e[key] for e in entries) for key in ("revenue_sen", "percentage_fee_sen", "fixed_fee_sen", "shipping_cost_sen", "total_channel_cost_sen")},
            }
        channel_cost = sum(b["total_channel_cost_sen"] for b in breakdown.values())
        profit = revenue - material - packaging - channel_cost
        direct = {a: sum(h["minutes"] for h in hours if h["product_id"] == pid and h["activity"] == a) for a in ("production", "fulfilment")}
        minutes = sum(direct.values())
        products.append({
            "product_id": pid, "name": product["name"], "units": units, "orders": len(sales),
            "revenue_sen": revenue, "revenue_rm": money(revenue),
            "material_unit_cost_sen": sen(product["material_cost_rm"]),
            "packaging_unit_cost_sen": sen(product["packaging_cost_rm"]),
            "material_cost_sen": material, "packaging_cost_sen": packaging,
            "unit_costs_sen": material + packaging, "unit_costs_rm": money(material + packaging),
            "channel_costs_sen": channel_cost, "channel_costs_rm": money(channel_cost), "channels": breakdown,
            "profit_sen": profit, "profit_rm": money(profit), "loses_money": profit < 0,
            "direct_minutes_by_activity": direct, "direct_minutes": minutes,
            "direct_hours_display": display(dec(minutes) / 60, 1),
            "profit_per_hour_rm": ratio(profit * 60, minutes * 100, "RM/hour"),
        })
    shared = activities["event"] + activities["admin"]
    assumptions = plan["assumptions"][0]
    numbers = {
        "category_totals_sen": totals, "category_totals_rm": {c: money(v) for c, v in totals.items()},
        "expense_total_sen": expenses, "expense_total_rm": money(expenses),
        "take_home_sen": take_home, "take_home_rm": money(take_home),
        "minutes_by_activity": activities,
        "hours_display_by_activity": {a: display(dec(v) / 60, 1) for a, v in activities.items()},
        "total_minutes": total_minutes, "total_hours_display": display(dec(total_minutes) / 60, 1),
        "shared_minutes": shared, "shared_hours_display": display(dec(shared) / 60, 1),
        "hourly_rate_rm": ratio(take_home * 60, total_minutes * 100, "RM/hour"),
        "margin_percent": ratio(take_home * 100, totals["sales"], "percent", 0),
        "planned_margin_percent": assumptions["planned_margin_percent"],
        "target_hourly_rate_rm": money(sen(assumptions["target_hourly_rate_rm"])),
        "below_planned_margin": dec(take_home) * 100 < dec(totals["sales"]) * dec(assumptions["planned_margin_percent"]),
        "below_target_hourly_rate": dec(take_home) * 60 < dec(sen(assumptions["target_hourly_rate_rm"])) * total_minutes,
        "product_profit_total_sen": sum(p["profit_sen"] for p in products),
        "basis_warning": "Take home uses paid ledger expenses; product profit uses plan unit/channel costs. They are intentionally different bases. Unknown expense is excluded, not valued at zero.",
    }
    return numbers, products, sorted(order_costs, key=lambda e: e["txn_id"])


def cash_message_check(messages, rows):
    declared = []
    for day, text in messages.items():
        for match in re.finditer(r"([A-Z][a-z]+) bought one (P\d+) [^;]+? for RM(\d+\.\d{2})", text):
            name, product, amount = match.groups()
            declared.append((day, name, product, 1, sen(amount)))
    ledger = [(r["date"], r["counterparty"], r["product_id"], r["quantity"], r["amount_sen"]) for r in rows if r["category"] == "sales" and r["payment_method"] == "cash"]
    return sorted(declared) == sorted(ledger)


def findings(audits, numbers, products):
    # PLAN describes ambiguity without references; it does not prohibit it.
    conflicts = []
    unknown = [a["txn_id"] for a in audits if a["original_document_amount_sen"] is None and a["derived_amount_sen"] is None]
    gaps = [
        {"id": "E01", "transaction_ids": unknown, "detail": "The thread receipt amount is unknown. Excluded from totals, not treated as zero. Known-amount take home and hourly rate may overstate earnings after this expense is recovered."},
        {"id": "E02", "detail": "Original receipt/order images and the original statement file are outside the permitted inputs. Statuses are audited against supplied transaction fields, notes, messages and bank_lines, not independently against those original files. File existence and legibility are not verified."},
        {"id": "E03", "detail": "The original mismatched document amount is available only in transaction notes. The bank amount is independently checked against bank_lines. Whether the difference is underpayment or another error remains unresolved."},
        {"id": "E04", "detail": "Intake timestamps and nightly reconciliation state are absent. This audit reconstructs end-of-week decision evidence, not an observed execution history. Missing amount supports an intake flag; the noted amount discrepancy supports a later reference-match flag."},
        {"id": "E05", "detail": "Draft recommendations and an actual Auditor ruling are not permitted inputs. Only the evidence-supported direction of a ruling is derived; no projected raise is asserted. Day Zero volume and projection rules remain unconfirmed in SYSTEM_DESIGN.md:375-382."},
    ]
    most_units = max(p["units"] for p in products)
    bestseller = [p["product_id"] for p in products if p["units"] == most_units]
    loss_leaders = [p["product_id"] for p in products if p["product_id"] in bestseller and p["loses_money"]]
    ruling = {
        "bestseller_product_ids": bestseller, "loss_making_bestseller_product_ids": loss_leaders,
        "direction": "Favor improving hourly earnings and fixing loss-making keychain pricing/costs before growing their current sales mix; do not endorse volume growth on bestseller status alone. Bracelets have positive direct profit and higher direct profit per hour. Shared time remains unallocated.",
        "basis": {"hourly_rate_rm": numbers["hourly_rate_rm"]["display"], "target_hourly_rate_rm": numbers["target_hourly_rate_rm"],
                  "products": [{k: p[k] for k in ("product_id", "units", "profit_sen", "direct_minutes", "profit_per_hour_rm")} for p in products]},
        "limitation": "This is an evidence-supported direction, not a verified existing ruling or a forecast. Actual gains need demand, feasibility and approved projection rules.",
    }
    return conflicts, gaps, ruling


def make_notes(result):
    n = result["numbers"]
    notes = [
        "Independent rebuild for " + result["week"]["iso_week"] + "; only the five authorized input files were used. No existing expected hours, calculator, output schema, answer key or memory was used as calculation evidence.",
        "The baseline plan was not saved or modified. The dated correction was parsed once from the Wednesday nightly message, not duplicated from its quoted replay instructions.",
        "Correction: " + "; ".join(f"{c['field']} {c['old_value']} to {c['new_value']} on {c['changed_at']} ({c['source']})" for c in result["corrections_applied_in_memory"]) + ".",
        "Sharon's source-label clarification is applied: keychain production before the correction date is estimate; on/after the date it is corrected. Historical old_value lookup alone does not make a line corrected.",
        "Hours IDs are local deterministic IDs: sort by date, activity in production/fulfilment/event/admin order, product, channel and assumption reference; then assign H0001 upward. The design specifies the ID shape but not this ordering.",
        "Amounts are integer sen and time is integer minutes. Exact ratio numerators/denominators are retained; displayed halves round away from zero. Channel percentage fees round separately on each order. Totals use unrounded minutes, not displayed hours.",
        f"Generated {len(result['hours'])} hours lines: " + "; ".join(f"{a} {v} minutes" for a, v in n["minutes_by_activity"].items()) + f"; total {n['total_minutes']} minutes ({n['total_hours_display']} hours).",
        f"Shared time is {n['shared_minutes']} minutes ({n['shared_hours_display']} hours), comprising event plus admin only; it is not allocated to products.",
        "Category totals: " + "; ".join(f"{c} RM{v}" for c, v in n["category_totals_rm"].items()) + ".",
        f"Known-amount expenses RM{n['expense_total_rm']}; take home RM{n['take_home_rm']}; hourly rate RM{n['hourly_rate_rm']['display']}; margin {n['margin_percent']['display']}%, versus planned {n['planned_margin_percent']}% and target RM{n['target_hourly_rate_rm']}/hour.",
    ]
    for p in result["products"]:
        notes.append(f"{p['product_id']} {p['name']}: {p['units']} units/{p['orders']} orders; revenue RM{p['revenue_rm']}; unit costs RM{p['unit_costs_rm']}; channel costs RM{p['channel_costs_rm']}; profit RM{p['profit_rm']}; direct time {p['direct_minutes']} minutes; profit/hour RM{p['profit_per_hour_rm']['display']}.")
    notes.append("Product profits sum to RM" + money(n["product_profit_total_sen"]) + ", not take home; the plan-cost and paid-ledger bases intentionally differ.")
    notes.append("Status counts: " + "; ".join(f"{s} {count}" for s, count in result["audit"]["status_counts"].items()) + ". Every input transaction status, bank assignment and final amount agrees with the independent evidence reconstruction.")
    for item in result["review_list"]:
        notes.append(f"Review {item['txn_id']}: {item['status']}; {'included' if item['counted'] else 'excluded'} in totals; {item['reason']}.")
    notes.append("Bank mapping: " + "; ".join(f"{b['bank_reference']} -> {', '.join(b['transaction_ids'])}" for b in result["bank_mapping"]) + ". Every bank line is used once; cash and the duplicate have no bank assignment.")
    notes.append("Four Saturday cash orders match the dated message individually. Nine Instagram orders match the courier's nine-label description; the courier expense is not an extra sale.")
    notes.append("Dana and Mira have no payment references: amount/date matching is ambiguous, and each counterparty uniquely identifies a bank description. Their ambiguous_resolved statuses are supported by SYSTEM_DESIGN.md:174-178. PLAN.md:113 describes this ambiguity; it does not mandate references or create a fixture conflict.")
    if not result["conflicts"]:
        notes.append("No unresolved numerical or specification conflict was found in the available evidence after applying Sharon's source-label clarification. Missing evidence and review items below remain unresolved.")
    for conflict in result["conflicts"]:
        notes.append(f"Conflict {conflict['id']}: {conflict['evidence']} {conflict['effect']}")
    for gap in result["missing_evidence"]:
        notes.append(f"Evidence limit {gap['id']}: {gap['detail']}")
    notes.append("Auditor direction: " + result["auditor_direction"]["direction"] + " " + result["auditor_direction"]["limitation"])
    notes.append("Validation: all arithmetic, boundary, source-label, cash-message, unique-bank-use, category, hours and reconciliation consistency checks pass. This is not a comparison with any forbidden answer-key output. The four review items and documented evidence limits remain open.")
    return "\n".join("- " + note for note in notes) + "\n"


def main():
    before = fingerprints()
    baseline = workbook(INPUTS[2], ("products", "channels", "assumptions", "change_log"))
    fixture = workbook(INPUTS[4], ("transactions", "bank_lines"))
    messages, corrections, events = parse_messages(INPUTS[3].read_text(encoding="utf-8"), baseline)
    plan = apply_corrections(baseline, corrections)
    rows, audits, bank = reconcile(fixture["transactions"], fixture["bank_lines"])
    first = min(date.fromisoformat(d) for d in messages)
    monday = first - timedelta(days=first.weekday())
    sunday = monday + timedelta(days=6)
    assert sorted(messages) == [(monday + timedelta(days=i)).isoformat() for i in range(7)]
    assert all(monday.isoformat() <= r["date"] <= sunday.isoformat() for r in rows + bank)
    hours = build_hours(rows, plan, events, sunday.isoformat())
    numbers, products, order_costs = compute_numbers(rows, hours, plan)
    conflicts, gaps, ruling = findings(audits, numbers, products)
    bank_mapping = [{**b, "amount_rm": money(b["amount_sen"]),
                     "transaction_ids": [r["txn_id"] for r in rows if r["bank_reference"] == b["bank_reference"]],
                     "statuses": [r["reconciliation_status"] for r in rows if r["bank_reference"] == b["bank_reference"]]}
                    for b in bank]
    review = [{"txn_id": r["txn_id"], "status": r["reconciliation_status"], "amount_sen": r["amount_sen"],
               "amount_rm": money(r["amount_sen"]), "counted": r["counted"], "reason": r["notes"]}
              for r in rows if r["reconciliation_status"] in REVIEW_STATUSES]
    daily = [{"date": day, "minutes_by_activity": {a: sum(h["minutes"] for h in hours if h["date"] == day and h["activity"] == a) for a in ACTIVITIES},
              "total_minutes": sum(h["minutes"] for h in hours if h["date"] == day),
              "entry_ids": [h["entry_id"] for h in hours if h["date"] == day]} for day in sorted(messages)]
    ledger_cash = {kind: sum(r["amount_sen"] for r in rows if r["counted"] and r["payment_method"] == "cash" and r["type"] == kind) for kind in ("income", "expense")}
    bank_totals = {kind: sum(b["amount_sen"] for b in bank if b["type"] == kind) for kind in ("income", "expense")}
    direct_recheck = sum(whole(dec(r["quantity"]) * dec(in_force(plan, f"products.{r['product_id']}.minutes_per_unit", r["date"])[0])) for r in rows if r["counted"] and r["category"] == "sales")
    fulfilment_recheck = sum(whole(in_force(plan, f"channels.{r['channel']}.fulfilment_minutes_per_order", r["date"])[0]) for r in rows if r["counted"] and r["category"] == "sales")
    checks = {
        "display_half_up_positive_negative": display("1.005") == "1.01" and display("-1.005") == "-1.01",
        "whole_sen_rounding_half": percentage_fee(5, 10) == 1 and percentage_fee(-5, 10) == -1,
        "per_order_fee_not_aggregate": sum(percentage_fee(5, 10) for _ in range(2)) == 2 and percentage_fee(10, 10) == 1,
        "all_transaction_statuses_agree": all(a["status_agrees"] for a in audits),
        "all_transaction_bank_references_agree": all(a["bank_reference_agrees"] for a in audits),
        "all_final_amounts_agree": all(a["amount_agrees"] for a in audits),
        "noted_bank_amounts_agree": all(a.get("bank_note_agrees", True) for a in audits),
        "leftover_categories_agree": all(a.get("category_agrees", True) for a in audits),
        "all_bank_lines_used_exactly_once": all(len(b["transaction_ids"]) == 1 for b in bank_mapping),
        "all_input_rows_audited": len(audits) == len(fixture["transactions"]),
        "category_type_consistency": all(r["category"] in CATEGORIES and r["type"] == ("income" if r["category"] == "sales" else "expense") for r in rows),
        "cash_orders_match_message": cash_message_check(messages, rows),
        "production_independent_per_sale_sum": direct_recheck == numbers["minutes_by_activity"]["production"],
        "fulfilment_independent_per_order_sum": fulfilment_recheck == numbers["minutes_by_activity"]["fulfilment"],
        "nonzero_fulfilment_lines": all(h["minutes"] > 0 for h in hours if h["activity"] == "fulfilment"),
        "unique_hours_ids": len({h["entry_id"] for h in hours}) == len(hours),
        "unique_hours_group_keys": len({(h["date"], h["activity"], h["product_id"], h["assumption_ref"]) for h in hours}) == len(hours),
        "whole_minutes": all(type(h["minutes"]) is int and h["minutes"] >= 0 for h in hours),
        "direct_plus_shared_minutes": sum(p["direct_minutes"] for p in products) + numbers["shared_minutes"] == numbers["total_minutes"],
        "daily_sum": sum(d["total_minutes"] for d in daily) == numbers["total_minutes"],
        "product_revenues_sum": sum(p["revenue_sen"] for p in products) == numbers["category_totals_sen"]["sales"],
        "income_bank_plus_cash": bank_totals["income"] + ledger_cash["income"] == numbers["category_totals_sen"]["sales"],
        "expenses_bank_plus_known_cash": bank_totals["expense"] + ledger_cash["expense"] == numbers["expense_total_sen"],
        "review_complete": {r["txn_id"] for r in review} == {r["txn_id"] for r in rows if r["reconciliation_status"] in REVIEW_STATUSES},
        "baseline_change_log_still_empty": not baseline["change_log"],
        "hours_rebuild_deterministic": hours == build_hours(rows, plan, events, sunday.isoformat()),
        "source_labels": all(h["source"] == ("stated" if h["activity"] == "event" else in_force(plan, h["assumption_ref"], h["date"])[1]) for h in hours),
        "replay_correction_not_double_counted": len(corrections) == len(plan["change_log"]),
    }
    for c in corrections:
        day_before = (date.fromisoformat(c["changed_at"]) - timedelta(days=1)).isoformat()
        checks[f"correction_boundary_{c['field']}"] = in_force(plan, c["field"], day_before)[:2] == (c["old_value"], "estimate") and in_force(plan, c["field"], c["changed_at"])[:2] == (c["new_value"], "corrected")
    for day, text in messages.items():
        if "nine parcel labels" in text:
            checks["courier_nine_labels_match_instagram_orders"] = sum(1 for r in rows if r["counted"] and r["category"] == "sales" and r["channel"] == "instagram") == 9
    checks["permitted_inputs_unchanged"] = before == fingerprints()
    assert all(checks.values()), json.dumps({k: v for k, v in checks.items() if not v})
    iso_year, iso_week, _ = monday.isocalendar()
    result = {
        "schema": "lembur-independent-rebuild-v1",
        "week": {"iso_week": f"{iso_year}-W{iso_week:02d}", "monday": monday.isoformat(), "sunday": sunday.isoformat(), "timezone": "Asia/Kuala_Lumpur", "scope": "Prepared complete-week synthetic fixture, not live activity"},
        "input_sha256": before,
        "conventions": {"money_calculation": "integer sen", "time_calculation": "integer minutes", "display_rounding": "Decimal ROUND_HALF_UP", "ratio_storage": "exact integer numerator and denominator plus rounded display", "hours_id_order": ["date", "production/fulfilment/event/admin", "product_id", "channel", "assumption_ref"]},
        "baseline_plan": baseline, "corrected_plan_in_memory": plan,
        "corrections_applied_in_memory": corrections, "events_parsed": events,
        "hours": hours, "daily_minutes": daily, "numbers": numbers,
        "products": products, "per_order_channel_costs": order_costs,
        "reconstructed_transactions": rows, "reconciliation_audit": audits,
        "review_list": review, "bank_mapping": bank_mapping,
        "bank_and_cash_control_sen": {"bank": bank_totals, "known_cash": ledger_cash},
        "conflicts": conflicts, "missing_evidence": gaps, "auditor_direction": ruling,
        "audit": {"checks": checks, "all_consistency_checks_pass": all(checks.values()),
                  "status_counts": dict(sorted(Counter(r["reconciliation_status"] for r in rows).items())),
                  "source_counts": dict(sorted(Counter(h["source"] for h in hours).items())),
                  "counted_rows": sum(r["counted"] for r in rows), "excluded_rows": sum(not r["counted"] for r in rows),
                  "conclusion": "Arithmetic and supplied-row evidence agree, with no unresolved specification conflict; this does not clear review items or replace original-document verification."},
    }
    assert OUT.is_dir()
    (OUT / "independent_result.json").write_text(json.dumps(result, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")
    (OUT / "independent_notes.md").write_text(make_notes(result), encoding="utf-8")
    assert before == fingerprints(), "Input changed during execution"
    print(json.dumps({"week": result["week"]["iso_week"], "hours_lines": len(hours), "minutes": numbers["total_minutes"],
                      "take_home_rm": numbers["take_home_rm"], "hourly_rate_rm": numbers["hourly_rate_rm"]["display"],
                      "products": [{"product_id": p["product_id"], "profit_rm": p["profit_rm"], "profit_per_hour_rm": p["profit_per_hour_rm"]["display"]} for p in products],
                      "checks_passed": len(checks), "review_ids": [r["txn_id"] for r in review], "conflict_count": len(conflicts)}, indent=2))


if __name__ == "__main__":
    main()

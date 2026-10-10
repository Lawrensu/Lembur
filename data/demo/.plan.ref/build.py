"""Build the approved, uncorrected Night Shift input plan."""
from pathlib import Path
try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[3]

def xl_color(css_hex):
    return "FF" + css_hex.removeprefix("#").upper()

XL_PRIMARY = xl_color("#245C73")
XL_WHITE = xl_color("#FFFFFF")
XL_DARK = xl_color("#172B36")
XL_PALE = xl_color("#EEF5F7")

def save_book(path, sheets, title):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.properties.title = title
    wb.properties.creator = "Lembur"
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
        ws.row_dimensions[1].height = 44
        for col, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(col)].width = min(50, max(16, len(header)+3))
            c = ws.cell(1, col)
            c.font = Font(name="Calibri", size=11, bold=True, color=XL_WHITE)
            c.fill = PatternFill("solid", fgColor=XL_PRIMARY)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            if header in ("source_file", "notes", "assumption_ref", "description", "reason", "field"):
                ws.column_dimensions[get_column_letter(col)].width = 48
            for r in range(2, len(rows)+2):
                c = ws.cell(r, col)
                c.font = Font(name="Calibri", size=11, color=XL_DARK)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if r % 2 == 0:
                    c.fill = PatternFill("solid", fgColor=XL_PALE)
                if header.endswith("_rm"):
                    c.number_format = "0.00"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)

if __name__ == "__main__":
    sheets = {
        "products": (["product_id", "name", "price_rm", "material_cost_rm", "packaging_cost_rm", "minutes_per_unit"], [
            ["P01", "Beaded keychain", 5, 2.5, 0.5, 20],
            ["P02", "Beaded bracelet", 18, 5, 1, 35]]),
        "channels": (["channel", "fee_percent", "fee_per_order_rm", "shipping_cost_per_order_rm", "fulfilment_minutes_per_order"], [
            ["instagram", 0, 0, 4, 5], ["telegram", 0, 0, 0, 3], ["bazaar", 0, 0, 0, 0]]),
        "assumptions": (["admin_hours_per_week", "planned_margin_percent", "target_hourly_rate_rm"], [[2, 40, 12]]),
        "change_log": (["changed_at", "field", "old_value", "new_value", "reason", "source"], [])
    }
    out = ROOT / "data/demo/plan.xlsx"
    save_book(out, sheets, "Lembur baseline plan - synthetic demo")
    check = openpyxl.load_workbook(out, data_only=True)
    assert check.sheetnames == list(sheets)
    for name, (headers, rows) in sheets.items():
        actual = list(check[name].values)
        assert list(actual[0]) == headers
        assert [list(r) for r in actual[1:]] == rows
    print(f"PASS: {out}; four sheets, baseline P01=20, change_log empty")
